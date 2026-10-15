import sys
import os
import pandas as pd
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
import urllib.parse
from sqlalchemy import create_engine

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass


def get_connection_string():
    """Stringa di connessione ODBC ad ARCA, letta dalle variabili d'ambiente (.env)."""
    conn_str = (
        "Driver={SQL Server};"
        f"Server={os.environ['SERVER_DB']};"
        f"Database={os.environ['DB_ARCA']};"
    )
    if os.getenv("DB_USER"):
        return conn_str + f"UID={os.environ['DB_USER']};PWD={os.getenv('DB_PASSWORD', '')};"
    return conn_str + "Trusted_Connection=yes;"


class ExtractionApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Estrattore Dati Ordini")
        self.root.geometry("600x520")
        self.root.resizable(False, False)

        # Stile e configurazione
        self.style = ttk.Style()
        self.style.theme_use('clam')
        self.style.configure('TButton', font=('Arial', 12), padding=5)
        self.style.configure('TLabel', font=('Arial', 12))
        self.style.configure('TEntry', font=('Arial', 12))
        self.style.configure('TCheckbutton', font=('Arial', 11))

        # Creazione dei frame per le diverse "pagine"
        self.page1_frame = ttk.Frame(root, padding="20")
        self.page2_frame = ttk.Frame(root, padding="20")

        # Dizionario per memorizzare le colonne e le relative checkbox
        self.column_vars = {}  # SPOSTATO QUI, PRIMA DELL'INIZIALIZZAZIONE DELLE PAGINE

        # Inizializzazione delle pagine
        self.setup_page1()
        self.setup_page2()

        # Mostra la prima pagina all'avvio
        self.show_page1() 
        
    def setup_page1(self):
        """Configura la prima pagina (parametri di ricerca)"""
        main_frame = self.page1_frame
        
        # Titolo
        title_label = ttk.Label(main_frame, text="Estrattore Dati Ordini", font=('Arial', 16, 'bold'))
        title_label.pack(pady=10)
        
        # Frame per i campi di input
        input_frame = ttk.Frame(main_frame)
        input_frame.pack(fill=tk.X, pady=10)
        
        # Anno Esercizio (obbligatorio, salvo ricerca per Identificativo)
        anno_frame = ttk.Frame(input_frame)
        anno_frame.pack(fill=tk.X, pady=5)
        ttk.Label(anno_frame, text="Anno Esercizio*:").pack(side=tk.LEFT, padx=5)
        self.anno_var = tk.StringVar()
        self.anno_entry = ttk.Entry(anno_frame, textvariable=self.anno_var, width=10)
        self.anno_entry.pack(side=tk.LEFT, padx=5)
        # current_year = datetime.now().year
        # self.anno_var.set(str(current_year))
        
        # Codice Documento (obbligatorio, salvo ricerca per Identificativo)
        doc_frame = ttk.Frame(input_frame)
        doc_frame.pack(fill=tk.X, pady=5)
        ttk.Label(doc_frame, text="Codice Documento*:").pack(side=tk.LEFT, padx=5)
        self.doc_var = tk.StringVar()
        self.doc_entry = ttk.Entry(doc_frame, textvariable=self.doc_var, width=10)
        self.doc_entry.pack(side=tk.LEFT, padx=5)
        # self.doc_var.set("FTV")
        
        # Numero Documento (obbligatorio, salvo ricerca per Identificativo)
        num_frame = ttk.Frame(input_frame)
        num_frame.pack(fill=tk.X, pady=5)
        ttk.Label(num_frame, text="Numero Documento*:").pack(side=tk.LEFT, padx=5)
        self.num_var = tk.StringVar()
        self.num_entry = ttk.Entry(num_frame, textvariable=self.num_var, width=20)
        self.num_entry.pack(side=tk.LEFT, padx=5)
        # Checkbox per la corrispondenza esatta
        self.exact_match_var = tk.BooleanVar(value=False)
        exact_match_cb = ttk.Checkbutton(num_frame, text="Esatta corrispondenza", variable=self.exact_match_var)
        exact_match_cb.pack(side=tk.LEFT, padx=5)
        
        # Codice Cliente (obbligatorio, salvo ricerca per Identificativo)
        cliente_frame = ttk.Frame(input_frame)
        cliente_frame.pack(fill=tk.X, pady=5)
        ttk.Label(cliente_frame, text="Codice Cliente*:").pack(side=tk.LEFT, padx=5)
        self.cliente_var = tk.StringVar()
        self.cliente_entry = ttk.Entry(cliente_frame, textvariable=self.cliente_var, width=20)
        self.cliente_entry.pack(side=tk.LEFT, padx=5)

        # Separatore visivo
        sep = ttk.Separator(input_frame, orient='horizontal')
        sep.pack(fill=tk.X, pady=10)

        # Identificativo (Facoltativo, ma se compilato basta da solo)
        id_frame = ttk.Frame(input_frame)
        id_frame.pack(fill=tk.X, pady=5)
        ttk.Label(id_frame, text="Identificativo (Id_DOTes):").pack(side=tk.LEFT, padx=5)
        self.id_var = tk.StringVar()
        self.id_entry = ttk.Entry(id_frame, textvariable=self.id_var, width=20)
        self.id_entry.pack(side=tk.LEFT, padx=5)

        id_help_label = ttk.Label(
            input_frame,
            text="Se compili l'Identificativo, la ricerca avviene solo su di esso\n"
                 "e gli altri campi sopra diventano ininfluenti.",
            font=('Arial', 9),
            foreground="#555555",
            justify=tk.LEFT
        )
        id_help_label.pack(anchor=tk.W, padx=5, pady=(0, 5))
        
        # Pulsante per passare alla pagina successiva
        next_btn = ttk.Button(main_frame, text="Avanti →", command=self.validate_and_show_page2)
        next_btn.pack(pady=15)
        
        # Etichetta campi obbligatori
        required_label = ttk.Label(main_frame, text="* Campi obbligatori se non si usa l'Identificativo", font=('Arial', 10))
        required_label.pack(anchor=tk.W, pady=5)
        
        # Stato connessione
        self.status_var = tk.StringVar()
        self.status_var.set("Pronto")
        status_label = ttk.Label(main_frame, textvariable=self.status_var, font=('Arial', 10))
        status_label.pack(side=tk.BOTTOM, fill=tk.X)
    
    def setup_page2(self):
        """Configura la seconda pagina (selezione colonne)"""
        main_frame = self.page2_frame
        
        # Titolo
        title_label = ttk.Label(main_frame, text="Seleziona le colonne da estrarre", font=('Arial', 16, 'bold'))
        title_label.pack(pady=10)
        
        # Frame per le colonne selezionabili
        columns_frame = ttk.LabelFrame(main_frame, text="Colonne disponibili")
        columns_frame.pack(fill=tk.BOTH, expand=True, pady=10, padx=5)
        
        # Definizione delle colonne disponibili
        # Nome visualizzato : (nome colonna SQL, selezionato di default)
        available_columns = {
            "Codice Articolo": ("DR.Cd_AR", True),
            "EAN PZ": ("AA.Alias", False),
            "Descrizione": ("DR.Descrizione", True),
            "Unità di Misura": ("DR.Cd_ARMisura", True),
            "Fattore di Conversione": ("FATTORE_CONVERSIONE", True),
            "Prezzo Unitario": ("PREZZO_UNITARIO", True),
            "Quantità": ("QUANTITA", True),
            "Prezzo Vendita Lordo": ("DR.PrezzoTotaleMovE", True),
            "Aliquota IVA": ("ALIQUOTA_IVA", True),
            "Prezzo Vendita Netto": ("PREZZO_VENDITA_NETTO", False),
            "Prezzo Vendita Consigliato": ("PREZZO_VENDITA_CONSIGLIATO", False)
        }
        
        # Creazione dei checkbox per le colonne
        left_col = ttk.Frame(columns_frame)
        left_col.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        right_col = ttk.Frame(columns_frame)
        right_col.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Dividi le colonne in due colonne visive
        half = len(available_columns) // 2
        items = list(available_columns.items())
        
        # Prima colonna di checkbox
        for i, (display_name, (sql_name, default)) in enumerate(items[:half]):
            var = tk.BooleanVar(value=default)
            self.column_vars[sql_name] = var
            cb = ttk.Checkbutton(left_col, text=display_name, variable=var)
            cb.pack(anchor=tk.W, pady=3)
        
        # Seconda colonna di checkbox
        for i, (display_name, (sql_name, default)) in enumerate(items[half:]):
            var = tk.BooleanVar(value=default)
            self.column_vars[sql_name] = var
            cb = ttk.Checkbutton(right_col, text=display_name, variable=var)
            cb.pack(anchor=tk.W, pady=3)
        
        # Frame per i pulsanti
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(pady=20)
        
        # Pulsante per tornare alla pagina precedente
        back_btn = ttk.Button(button_frame, text="← Indietro", command=self.show_page1)
        back_btn.pack(side=tk.LEFT, padx=10)
        
        # Pulsante di estrazione
        extract_btn = ttk.Button(button_frame, text="Estrai Dati", command=self.extract_data)
        extract_btn.pack(side=tk.LEFT, padx=10)
        
        # Pulsante di chiusura
        close_btn = ttk.Button(button_frame, text="Chiudi", command=self.root.quit)
        close_btn.pack(side=tk.LEFT, padx=10)
    
    def show_page1(self):
        """Mostra la prima pagina e nasconde la seconda"""
        self.page2_frame.pack_forget()
        self.page1_frame.pack(fill=tk.BOTH, expand=True)
    
    def show_page2(self):
        """Mostra la seconda pagina e nasconde la prima"""
        self.page1_frame.pack_forget()
        self.page2_frame.pack(fill=tk.BOTH, expand=True)
    
    def validate_and_show_page2(self):
        """Valida i campi della prima pagina e mostra la seconda se tutto è corretto"""
        if self.validate_inputs():
            self.show_page2()

    def is_using_identificativo(self):
        """Restituisce True se l'utente ha compilato l'Identificativo (Id_DOTes)"""
        return bool(self.id_var.get().strip())
    
    def validate_inputs(self):
        """Valida i campi di input. Se l'Identificativo è compilato, basta da solo."""
        identificativo = self.id_var.get().strip()

        # --- Modalità ricerca per Identificativo (Id_DOTes) ---
        if identificativo:
            if not identificativo.isdigit():
                messagebox.showerror("Errore", "L'Identificativo (Id_DOTes) deve essere un valore numerico")
                return False
            # Con l'Identificativo valorizzato, tutti gli altri campi sono ininfluenti
            return True

        # --- Modalità ricerca "classica" tramite Anno/Documento/Numero/Cliente ---
        anno = self.anno_var.get().strip()
        doc = self.doc_var.get().strip()
        num = self.num_var.get().strip()
        cliente = self.cliente_var.get().strip()

        if not anno:
            messagebox.showerror("Errore", "L'anno di esercizio è obbligatorio")
            return False
        
        if not anno.isdigit():
            messagebox.showerror("Errore", "L'anno di esercizio deve essere un valore numerico")
            return False
            
        if not doc:
            messagebox.showerror("Errore", "Il codice documento è obbligatorio")
            return False
            
        if not num:
            messagebox.showerror("Errore", "Il numero documento è obbligatorio")
            return False
        
        if not cliente:
            messagebox.showerror("Errore", "Il codice cliente è obbligatorio")
            return False
            
        return True
    
    def get_connection(self):
        """Crea la connessione al database utilizzando SQLAlchemy"""
        params = urllib.parse.quote_plus(get_connection_string())
        engine = create_engine("mssql+pyodbc:///?odbc_connect=%s" % params)
        return engine
    
    def get_selected_columns(self):
        """Restituisce l'elenco delle colonne selezionate"""
        selected = []
        
        # Controllo delle colonne selezionate
        for sql_name, var in self.column_vars.items():
            if var.get():
                selected.append(sql_name)
        
        return selected
    
    def get_column_expressions(self):
        """Restituisce le espressioni SQL per le colonne calcolate"""
        expressions = {
            "FATTORE_CONVERSIONE": """
                CASE 
                    WHEN DR.FattoreToUM1 = FLOOR(DR.FattoreToUM1) THEN CAST(FLOOR(DR.FattoreToUM1) AS VARCHAR)
                    ELSE CAST(DR.FattoreToUM1 AS VARCHAR)
                END
            """,
            "PREZZO_UNITARIO": """
                CAST(CONVERT(DECIMAL(18,2), ROUND(DR.PrezzoUnitarioV, 2)) AS VARCHAR)
            """,
            "QUANTITA": """
                CASE 
                    WHEN DR.Qta = FLOOR(DR.Qta) THEN CAST(FLOOR(DR.Qta) AS VARCHAR)
                    ELSE CAST(DR.Qta AS VARCHAR)
                END
            """,
            "ALIQUOTA_IVA": """
                CASE 
                    WHEN DR.Cd_Aliquota = 'W22' THEN '22%'
                    WHEN DR.Cd_Aliquota = 'W10' THEN '10%'
                    ELSE DR.Cd_Aliquota
                END
            """,
            "PREZZO_VENDITA_NETTO": """
                CASE 
                    WHEN DR.Cd_Aliquota = 'W22' THEN 
                        CAST(CONVERT(DECIMAL(18,2), ROUND(DR.PrezzoTotaleMovE / 1.22, 2)) AS VARCHAR)
                    WHEN DR.Cd_Aliquota = 'W10' THEN 
                        CAST(CONVERT(DECIMAL(18,2), ROUND(DR.PrezzoTotaleMovE / 1.10, 2)) AS VARCHAR)
                    ELSE 
                        CAST(CONVERT(DECIMAL(18,2), ROUND(DR.PrezzoTotaleMovE, 2)) AS VARCHAR)
                END
            """,
            "PREZZO_VENDITA_CONSIGLIATO": """
            CAST(CONVERT(DECIMAL(18,2), ROUND((SELECT LSA.Prezzo FROM [dbo].[LSArticolo] LSA WHERE LSA.Cd_AR = DR.Cd_AR AND LSA.Id_LSRevisione = '1'), 2)) AS VARCHAR)
            """
        }
        return expressions

    def build_where_clauses(self, anno, doc, num, cliente, identificativo=None):
        """
        Costruisce le condizioni WHERE.
        Se identificativo è valorizzato, la ricerca avviene esclusivamente
        su DR.Id_DOTes, ignorando anno/doc/num/cliente.
        """
        if identificativo:
            return [
                "DR.Cd_AR IS NOT NULL",
                f"DR.[Id_DOTes] = {identificativo}"
            ]

        where_clauses = [
            "DR.Cd_AR IS NOT NULL",
            f"DR.[Cd_MGEsercizio] = {anno}",
            f"DR.[Cd_DO] = '{doc}'",
            f"DR.[Cd_CF] = '{cliente}'"
        ]

        if "%" in num:
            where_clauses.append(f"DR.[NumeroDoc] LIKE '{num}'")
        else:
            where_clauses.append(f"DR.[NumeroDoc] LIKE '%{num}%'")

        return where_clauses
    
    def build_dynamic_query(self):
        """Costruisce la query SQL dinamicamente basata sulle colonne selezionate (con eventuale filtro esatto)"""
        anno = self.anno_var.get().strip()
        doc = self.doc_var.get().strip()
        num = self.num_var.get().strip()
        cliente = self.cliente_var.get().strip()
        exact_match = self.exact_match_var.get()
        identificativo = self.id_var.get().strip()

        # Log per debug
        print(f"Costruzione query con: anno={anno}, doc={doc}, num={num}, cliente={cliente}, "
              f"esatta={exact_match}, identificativo={identificativo}")

        # Ottieni le colonne selezionate
        selected_columns = self.get_selected_columns()
        
        # Se non è selezionata alcuna colonna, mostra un errore
        if not selected_columns:
            messagebox.showerror("Errore", "Devi selezionare almeno una colonna da estrarre")
            return None
        
        # Definisci i nomi per la visualizzazione
        column_display_names = {
            "DR.Cd_AR": "CODICE_ARTICOLO",
            "AA.Alias": "EAN_PZ",
            "DR.Descrizione": "DESCRIZIONE",
            "DR.Cd_ARMisura": "ARMISURA",
            "FATTORE_CONVERSIONE": "FATTORE_CONVERSIONE",
            "PREZZO_UNITARIO": "PREZZO_UNITARIO",
            "QUANTITA": "QUANTITA",
            "DR.PrezzoTotaleMovE": "PREZZO_VENDITA_LORDO",
            "ALIQUOTA_IVA": "ALIQUOTA_IVA",
            "PREZZO_VENDITA_NETTO": "PREZZO_VENDITA_NETTO",
            "PREZZO_VENDITA_CONSIGLIATO": "PREZZO_VENDITA_CONSIGLIATO"
        }
        
        # Ottieni le espressioni per le colonne calcolate
        expressions = self.get_column_expressions()
        
        # Costruisci la lista delle colonne SQL
        select_parts = ["DR.Id_DOTes", "DR.NumeroDoc", "CONVERT(VARCHAR, DR.DataDoc, 103) AS DataDoc"]  # Colonne di servizio sempre incluse
        
        for col in selected_columns:
            if col in expressions:
                select_parts.append(f"{expressions[col]} AS {column_display_names[col]}")
            else:
                select_parts.append(f"{col} AS {column_display_names[col]}")
        
        select_clause = ", ".join(select_parts)

        # Se si usa l'Identificativo, il filtro esatto sul numero doc non ha senso: lo ignoriamo
        if identificativo:
            where_clauses = self.build_where_clauses(anno, doc, num, cliente, identificativo=identificativo)
        elif exact_match:
            where_clauses = self.build_where_clauses(anno, doc, num, cliente)
            # Sostituisce l'ultima clausola (LIKE) con l'uguaglianza esatta
            where_clauses[-1] = f"DR.[NumeroDoc] = '{num}'"
        else:
            where_clauses = self.build_where_clauses(anno, doc, num, cliente)

        where_clause = " AND ".join(where_clauses)
        
        query = f"""
        SELECT 
            {select_clause}
        FROM 
            [dbo].[DORig] DR
        LEFT JOIN 
            [dbo].[ARAlias] AA ON DR.Cd_AR = AA.Cd_AR AND AA.TipoAlias = 'E'
        WHERE 
            {where_clause}
        """
        
        return query
    
    def build_dynamic_query_without_exact_match(self, anno, doc, num, cliente, selected_columns, identificativo=None):
        """
        Costruisce la query senza applicare il filtro esatto sul numero documento
        (il filtro esatto, quando richiesto, viene applicato in Python su NumeroDoc).
        Se identificativo è valorizzato, la ricerca avviene solo su DR.Id_DOTes.
        """
        column_display_names = {
            "DR.Cd_AR": "CODICE_ARTICOLO",
            "AA.Alias": "EAN",
            "DR.Descrizione": "DESCRIZIONE",
            "DR.Cd_ARMisura": "ARMISURA",
            "FATTORE_CONVERSIONE": "FATTORE_CONVERSIONE",
            "PREZZO_UNITARIO": "PREZZO_UNITARIO",
            "QUANTITA": "QUANTITA",
            "DR.PrezzoTotaleMovE": "PREZZO_VENDITA_LORDO",
            "ALIQUOTA_IVA": "ALIQUOTA_IVA",
            "PREZZO_VENDITA_NETTO": "PREZZO_VENDITA_NETTO",
            "PREZZO_VENDITA_CONSIGLIATO": "PREZZO_VENDITA_CONSIGLIATO"
        }
        
        expressions = self.get_column_expressions()
        
        select_parts = ["DR.Id_DOTes", "DR.NumeroDoc", "CONVERT(VARCHAR, DR.DataDoc, 103) AS DataDoc"]
        
        for col in selected_columns:
            if col in expressions:
                select_parts.append(f"{expressions[col]} AS {column_display_names[col]}")
            else:
                select_parts.append(f"{col} AS {column_display_names[col]}")
        
        select_clause = ", ".join(select_parts)

        where_clauses = self.build_where_clauses(anno, doc, num, cliente, identificativo=identificativo)
        where_clause = " AND ".join(where_clauses)
        
        query = f"""
        SELECT 
            {select_clause}
        FROM 
            [dbo].[DORig] DR
        LEFT JOIN 
            [dbo].[ARAlias] AA ON DR.Cd_AR = AA.Cd_AR AND AA.TipoAlias = 'E'
        WHERE 
            {where_clause}
        """
        
        return query

    def show_multiple_docs_warning(self, df):
        """Mostra un avviso se ci sono più documenti diversi nei risultati"""
        try:
            unique_docs = df[['Id_DOTes', 'NumeroDoc', 'DataDoc']].drop_duplicates().reset_index(drop=True)

            if len(unique_docs) > 1:
                message = "Attenzione: i risultati provengono da documenti diversi:\n\n"

                for _, row in unique_docs.iterrows():
                    doc_date = row['DataDoc']
                    message += f"ID: {row['Id_DOTes']}, Numero: {row['NumeroDoc']}, Data: {doc_date}\n"

                messagebox.showwarning("Documenti multipli", message)

            return unique_docs
        except Exception as e:
            print(f"Errore nel controllo documenti multipli: {str(e)}")
            return df
    
    def extract_data(self):
        """Esegue l'estrazione dei dati e li salva in Excel"""
        if not self.validate_inputs():
            return

        try:
            self.status_var.set("Costruzione query...")
            self.root.update_idletasks()

            anno = self.anno_var.get().strip()
            doc = self.doc_var.get().strip()
            num = self.num_var.get().strip()
            cliente = self.cliente_var.get().strip()
            exact_match = self.exact_match_var.get()
            identificativo = self.id_var.get().strip()
            use_identificativo = bool(identificativo)

            selected_columns = self.get_selected_columns()

            if not selected_columns:
                messagebox.showerror("Errore", "Devi selezionare almeno una colonna da estrarre")
                return

            # Se è valorizzato l'Identificativo, la query filtra solo su Id_DOTes
            query = self.build_dynamic_query_without_exact_match(
                anno, doc, num, cliente, selected_columns,
                identificativo=identificativo if use_identificativo else None
            )

            self.status_var.set("Connessione al database...")
            self.root.update_idletasks()

            engine = self.get_connection()

            self.status_var.set("Esecuzione query...")
            self.root.update_idletasks()

            df = pd.read_sql(query, engine)

            if df.empty:
                if use_identificativo:
                    messagebox.showinfo("Informazione", f"Nessun documento trovato con Identificativo '{identificativo}'.")
                else:
                    messagebox.showinfo("Informazione", "Nessun risultato trovato con i criteri specificati.")
                self.status_var.set("Pronto")
                return

            # Il filtro esatto sul NumeroDoc ha senso solo nella ricerca "classica"
            if exact_match and not use_identificativo:
                self.status_var.set("Applicazione filtro esatto...")
                self.root.update_idletasks()

                def normalize_num(num_str):
                    try:
                        clean_str = num_str.strip()
                        return int(clean_str) if clean_str.isdigit() else clean_str
                    except:
                        return num_str

                target_num = normalize_num(num)

                df['NumeroDoc_Normalized'] = df['NumeroDoc'].apply(normalize_num)
                df_filtered = df[df['NumeroDoc_Normalized'] == target_num]

                if df_filtered.empty:
                    messagebox.showinfo("Informazione", f"Nessun documento con numero esattamente '{num}' trovato.")
                    self.status_var.set("Pronto")
                    return

                df = df_filtered.drop(columns=['NumeroDoc_Normalized'])

            unique_docs = self.show_multiple_docs_warning(df)

            export_df = df.drop(['Id_DOTes', 'NumeroDoc', 'DataDoc'], axis=1)
            
            self.status_var.set("Salvataggio dati...")
            self.root.update_idletasks()
            
            # Costruisci un nome file sensato in entrambe le modalità
            if use_identificativo:
                default_filename = f"Estratto_ID{identificativo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            else:
                num_clean = num.replace("%", "")
                default_filename = f"Estratto_{doc}_{num_clean}_{cliente}_{anno}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

            file_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile=default_filename
            )
            
            if not file_path:
                self.status_var.set("Operazione annullata")
                return
            
            export_df.to_excel(file_path, index=False, sheet_name="Dati Ordine")
            
            try:
                import openpyxl
                from openpyxl.styles import Font, Alignment, PatternFill
                
                wb = openpyxl.load_workbook(file_path)
                ws = wb.active
                
                header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")
                
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center')
                
                for column in ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    
                    for cell in column:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    
                    adjusted_width = max(max_length + 2, 12)
                    ws.column_dimensions[column_letter].width = adjusted_width
                
                wb.save(file_path)
                
            except ImportError:
                pass
            
            self.status_var.set(f"Dati estratti e salvati in: {file_path}")
            messagebox.showinfo("Successo", f"Estrazione completata con successo!\nFile salvato: {file_path}")
            
            try:
                os.startfile(file_path)
            except:
                pass
                
        except Exception as e:
            messagebox.showerror("Errore", f"Si è verificato un errore durante l'estrazione dei dati:\n{str(e)}")
            self.status_var.set("Errore")

# Funzione principale per avviare l'applicazione
def main():
    root = tk.Tk()
    app = ExtractionApp(root)
    root.mainloop()

if __name__ == "__main__":
    main()